"""Generate the synthetic demo corpus for the fictional company "Northwind Logistics".

Every person, place and figure is invented. Run from the repo root:

    cd backend && uv run python ../demo-corpus/generate.py

Output goes to demo-corpus/files/<collection>/ plus demo-corpus/manifest.json. The generated
files are committed so the demo works without running this script.
"""

import json
import random
import shutil
import sys
from pathlib import Path

import pymupdf
from docx import Document as DocxDocument

sys.path.insert(0, str(Path(__file__).parent))
from content_core import CORE  # noqa: E402

ROOT = Path(__file__).parent
OUT = ROOT / "files"
SEED = 20250101

COLLECTIONS = {
    "company": {
        "description": "Company-wide policies, procedures, depot handbooks and manuals.",
        "acl_groups": ["HR", "Engineering"],
    },
    "hr": {
        "description": "HR-only material: compensation, discipline, recruitment.",
        "acl_groups": ["HR"],
    },
    "engineering": {
        "description": "Engineering runbooks, standards and architecture.",
        "acl_groups": ["Engineering"],
    },
}

Section = tuple[str, list[str]]

# ---------------------------------------------------------------- writers


def _wrap(text: str, font: str, size: float, width: float) -> list[str]:
    lines: list[str] = []
    current = ""
    for word in text.split():
        trial = f"{current} {word}".strip()
        if current and pymupdf.get_text_length(trial, fontname=font, fontsize=size) > width:
            lines.append(current)
            current = word
        else:
            current = trial
    if current:
        lines.append(current)
    return lines


def write_pdf(path: Path, title: str, sections: list[Section], per_section_page: bool) -> None:
    doc = pymupdf.open()
    left, right, top, bottom = 72.0, 523.0, 90.0, 770.0
    width = right - left
    state = {"page": None, "y": top}

    def new_page() -> None:
        page = doc.new_page(width=595, height=842)
        n = doc.page_count
        page.insert_text((left, 40), "Northwind Logistics - Internal Use", fontsize=9, fontname="helv")
        page.insert_text((left, 810), f"Page {n}", fontsize=9, fontname="helv")
        state["page"], state["y"] = page, top

    def put(text: str, font: str, size: float, gap_after: float) -> None:
        for line in _wrap(text, font, size, width):
            if state["y"] > bottom:
                new_page()
            state["page"].insert_text((left, state["y"]), line, fontsize=size, fontname=font)  # type: ignore[attr-defined]
            state["y"] += size * 1.4
        state["y"] += gap_after

    new_page()
    put(title, "hebo", 22, 14)
    for i, (heading, paragraphs) in enumerate(sections):
        if per_section_page and i > 0:
            new_page()
        elif state["y"] > bottom - 80:
            new_page()
        put(heading, "hebo", 15, 6)
        for p in paragraphs:
            put(p, "helv", 11, 8)
    doc.set_metadata({"title": title, "creationDate": "", "modDate": "", "producer": "", "creator": ""})
    doc.save(path, deflate=True, garbage=4)
    doc.close()


def write_scanned_pdf(path: Path, title: str, sections: list[Section]) -> None:
    """Image-only PDF (no text layer), like a scanned paper document."""
    tmp = path.with_suffix(".tmp.pdf")
    write_pdf(tmp, title, sections, per_section_page=False)
    src = pymupdf.open(tmp)
    out = pymupdf.open()
    for page in src:
        pix = page.get_pixmap(dpi=80)
        new = out.new_page(width=page.rect.width, height=page.rect.height)
        new.insert_image(new.rect, pixmap=pix)
    out.set_metadata({"creationDate": "", "modDate": "", "producer": "", "creator": ""})
    out.save(path, deflate=True, garbage=4)
    out.close()
    src.close()
    tmp.unlink()


def write_docx(path: Path, title: str, sections: list[Section]) -> None:
    doc = DocxDocument()
    doc.core_properties.title = title
    doc.core_properties.author = "Northwind Logistics"
    doc.add_heading(title, level=0)
    for heading, paragraphs in sections:
        doc.add_heading(heading, level=1)
        for p in paragraphs:
            doc.add_paragraph(p)
    doc.save(path)


def write_html(path: Path, title: str, sections: list[Section]) -> None:
    body = "".join(
        f"<h2>{h}</h2>" + "".join(f"<p>{p}</p>" for p in ps) for h, ps in sections
    )
    path.write_text(
        f"<!doctype html><html><head><title>{title}</title></head><body>"
        "<nav>Home | Documents | Contact</nav>"
        f"<article><h1>{title}</h1>{body}</article>"
        "<footer>Northwind Logistics - internal</footer></body></html>",
        encoding="utf-8",
    )


def write_md(path: Path, title: str, sections: list[Section]) -> None:
    parts = [f"# {title}\n"]
    for h, ps in sections:
        parts.append(f"## {h}\n")
        parts.extend(f"{p}\n" for p in ps)
    path.write_text("\n".join(parts), encoding="utf-8")


def write_txt(path: Path, title: str, sections: list[Section]) -> None:
    parts = [title.upper(), ""]
    for h, ps in sections:
        parts.append(h)
        parts.append("")
        parts.extend(f"{p}\n" for p in ps)
    path.write_text("\n".join(parts), encoding="utf-8")


# ------------------------------------------------- procedural documents

FIRST = ["Amara", "Bruno", "Chiara", "Dmitri", "Elena", "Farid", "Greta", "Hana", "Ivo", "Jonas",
         "Keiko", "Lucía", "Mateo", "Nadia", "Oskar", "Petra", "Quentin", "Rosa", "Stefan", "Talia"]  # fmt: skip
LAST = ["Aldenmoor", "Brightwater", "Castellane", "Dunmore", "Everhart", "Fairbanks", "Grimsby",
        "Holloway", "Ivanov", "Jarrow", "Kestrel", "Lindqvist", "Marlowe", "Novak", "Okafor",
        "Pemberton", "Quill", "Rasmussen", "Sandoval", "Thorne"]  # fmt: skip
CITIES = ["Marlow Bay", "Harrowgate", "Eastfold", "Northmere", "Kingsreach", "Stonebridge",
          "Willowdene", "Ashcombe", "Redhaven", "Ferncliff", "Oakhurst", "Silverton", "Brackwater",
          "Thornbury", "Lakeside Cross", "Copperfield", "Windmere", "Highgate Vale", "Saltmarsh",
          "Dovedale"]  # fmt: skip


def depot_docs(rng: random.Random) -> list[tuple[str, str, str, str, str, bool, list[Section]]]:
    docs = []
    for i, city in enumerate(CITIES):
        manager = f"{FIRST[i]} {LAST[(i * 7) % len(LAST)]}"
        docks = rng.randint(6, 30)
        opens, closes = rng.choice([("05:00", "21:00"), ("06:00", "22:00"), ("07:00", "19:00")])
        cold = rng.choice([0, 120, 250, 400, 800])
        parking = rng.randint(20, 90)
        sections: list[Section] = [
            ("Site overview", [
                f"The {city} depot is managed by {manager}. It has {docks} loading dock doors and parking for {parking} trailers.",
                f"The depot is open from {opens} to {closes} on weekdays. Weekend dispatch runs with a reduced crew.",
            ]),
            ("Storage", [
                (f"The depot has a temperature-controlled area with a capacity of {cold} pallets."
                 if cold else "The depot has no temperature-controlled storage. Chilled freight is routed to another site."),
                "Dangerous goods are not accepted unless the shipment was pre-approved by the Safety Office.",
            ]),
            ("Emergency contacts", [
                f"In an emergency call the duty manager on extension {4000 + i}. The nearest muster point is the north car park.",
            ]),
        ]
        fmt = ["pdf", "docx", "html", "md"][i % 4]
        docs.append((f"depot-handbook-{city.lower().replace(' ', '-')}", "company", fmt,
                     f"{city} Depot Handbook", "2025-02-01", False, sections))
    return docs


VEHICLES = ["NL-Courier-S", "NL-Courier-M", "NL-Courier-L", "NL-Hauler-12", "NL-Hauler-18",
            "NL-Hauler-26", "NL-ChillVan", "NL-ChillTruck", "NL-Forklift-E2", "NL-Forklift-E4",
            "NL-Pallet-Jack", "NL-Yard-Tug", "NL-Cargo-Bike", "NL-Mini-Van", "NL-Tipper-8"]  # fmt: skip


def vehicle_docs(rng: random.Random) -> list[tuple[str, str, str, str, str, bool, list[Section]]]:
    docs = []
    for i, name in enumerate(VEHICLES):
        payload = rng.choice([180, 650, 1200, 3500, 8000, 12000, 18000])
        service = rng.choice([8000, 10000, 15000, 20000, 30000])
        tyre = rng.choice([2.2, 3.0, 5.5, 6.8, 8.3])
        battery = rng.choice([24, 48, 72, 96, 150])
        sections: list[Section] = [
            ("Specifications", [
                f"The {name} has a maximum payload of {payload} kg. The battery capacity is {battery} kWh.",
            ]),
            ("Maintenance", [
                f"The {name} must be serviced every {service} km or every 12 months, whichever comes first.",
                f"The recommended tyre pressure is {tyre} bar. Check tyre pressure before every shift.",
            ]),
            ("Safety", [
                f"Operators must hold a valid licence for the {name} class and complete a pre-use inspection. Any defect must be reported to the workshop before the vehicle is used.",
            ]),
        ]
        docs.append((f"manual-{name.lower()}", "company", "pdf" if i % 2 == 0 else "docx",
                     f"{name} Operator Manual", "2024-10-15", False, sections))
    return docs


SOP_TOPICS = [
    "Inbound goods receipt", "Outbound loading", "Damaged parcel handling", "Returns processing",
    "Pallet labelling", "Cold chain monitoring", "Hazardous goods screening", "Driver handover",
    "Proof of delivery", "Customs paperwork", "Cycle counting", "Forklift pre-use check",
    "Trailer sealing", "Missed delivery follow-up", "Overweight load handling", "Night shift start-up",
    "Fire evacuation drill", "Spill response", "Visitor sign-in", "Lost parcel investigation",
    "Peak season staffing",
]  # fmt: skip


def sop_docs(rng: random.Random) -> list[tuple[str, str, str, str, str, bool, list[Section]]]:
    docs = []
    for i, topic in enumerate(SOP_TOPICS):
        steps = rng.randint(4, 9)
        minutes = rng.choice([10, 15, 20, 30, 45])
        approver = rng.choice(["the shift supervisor", "the depot manager", "the quality officer"])
        sections: list[Section] = [
            ("Purpose", [f"This procedure describes how Northwind Logistics staff carry out {topic.lower()} consistently and safely."]),
            ("Procedure", [
                f"The procedure has {steps} steps and should take no longer than {minutes} minutes per occurrence.",
                f"Each completed occurrence is signed off by {approver}. Records are kept for 24 months.",
            ]),
            ("Escalation", [
                f"If the procedure cannot be completed within {minutes * 2} minutes the shift supervisor must be informed and the delay logged.",
            ]),
        ]
        docs.append((f"sop-{topic.lower().replace(' ', '-')}", "company", "docx" if i % 3 else "pdf",
                     f"SOP: {topic}", "2024-08-01", False, sections))
    return docs


EQUIPMENT = ["Handheld scanner HS-40", "Label printer LP-200", "Dock leveller DL-9", "Pallet wrapper PW-3",
             "Conveyor CV-12", "Weighbridge WB-50", "Barcode gate BG-7", "Rack system RS-4",
             "Temperature logger TL-1", "Forklift charger FC-22", "Safety light curtain SL-5",
             "Parcel sorter PS-8"]  # fmt: skip


def equipment_docs(rng: random.Random) -> list[tuple[str, str, str, str, str, bool, list[Section]]]:
    docs = []
    for name in EQUIPMENT:
        warranty = rng.choice([12, 24, 36, 60])
        inspect = rng.choice([1, 3, 6, 12])
        rating = rng.choice([110, 230, 400])
        sections: list[Section] = [
            ("Overview", [f"The {name} is standard depot equipment. It runs on a {rating} volt supply."]),
            ("Care", [
                f"Inspect the {name} every {inspect} months and record the result in the asset register. The manufacturer warranty lasts {warranty} months.",
            ]),
        ]
        docs.append((f"equipment-{name.lower().replace(' ', '-')}", "company", "html",
                     f"Equipment Sheet: {name}", "2024-06-01", False, sections))
    return docs


SCANNED = [
    ("scanned-delivery-note-1041", "Delivery Note 1041", [
        ("Delivery details", ["Delivered 14 pallets of mixed goods to the Redhaven depot. Received by the dock clerk and signed."]),
    ]),
    ("scanned-damage-report-207", "Damage Report 207", [
        ("Damage details", ["Two pallets arrived with crushed corners. Photographs were taken and the carrier was notified the same day."]),
    ]),
]  # fmt: skip

SHIPPING_RATES = ("shipping-rates-summary", "company", "md", "Shipping Rates Summary", "2025-01-10", False, [
    ("Standard parcels", ["Standard parcel delivery costs 4.90 euros for parcels up to 5 kg and 8.50 euros for parcels up to 20 kg."]),
    ("Pallets", ["Pallet delivery within the region costs 62 euros per pallet. Each additional pallet on the same route costs 48 euros."]),
])  # fmt: skip


def main() -> None:
    rng = random.Random(SEED)
    if OUT.exists():
        shutil.rmtree(OUT)
    for c in COLLECTIONS:
        (OUT / c).mkdir(parents=True)

    all_docs = [
        *CORE,
        *depot_docs(rng),
        *vehicle_docs(rng),
        *sop_docs(rng),
        *equipment_docs(rng),
        SHIPPING_RATES,
    ]
    manifest_docs: list[dict[str, str]] = []
    for slug, collection, fmt, title, modified, per_page, sections in all_docs:
        path = OUT / collection / f"{slug}.{fmt}"
        if fmt == "pdf":
            write_pdf(path, title, sections, per_page)
        elif fmt == "docx":
            write_docx(path, title, sections)
        elif fmt == "html":
            write_html(path, title, sections)
        elif fmt == "md":
            write_md(path, title, sections)
        else:
            write_txt(path, title, sections)
        manifest_docs.append({"path": str(path.relative_to(ROOT)), "collection": collection,
                              "modified_at": modified})
    for slug, title, sections in SCANNED:
        path = OUT / "company" / f"{slug}.pdf"
        write_scanned_pdf(path, title, sections)
        manifest_docs.append({"path": str(path.relative_to(ROOT)), "collection": "company",
                              "modified_at": "2025-01-20"})

    manifest = {
        "company": "Northwind Logistics (fictional)",
        "collections": [{"name": n, **c} for n, c in COLLECTIONS.items()],
        "documents": sorted(manifest_docs, key=lambda d: d["path"]),
    }
    (ROOT / "manifest.json").write_text(json.dumps(manifest, indent=2) + "\n", encoding="utf-8")
    print(f"wrote {len(manifest_docs)} documents")


if __name__ == "__main__":
    main()
