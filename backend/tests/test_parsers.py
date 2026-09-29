import pytest

from app.config import Settings
from app.ingestion.parsers import parse_document
from app.ingestion.parsers.docx import parse_docx
from app.ingestion.parsers.html import parse_html
from app.ingestion.parsers.pdf import parse_pdf
from app.ingestion.parsers.text import parse_text
from app.ingestion.types import MIME_PDF, ParseError
from tests.support.builders import (
    BODY,
    make_docx,
    make_encrypted_pdf,
    make_pdf,
    make_scanned_pdf,
    rewrite_zip,
)

S = Settings()


# ---------------------------------------------------------------- PDF
def test_pdf_detects_headings_by_font_size_and_keeps_page_numbers() -> None:
    data = make_pdf(
        [
            [("Annual Leave", 20), (BODY, 11), ("Carry forward", 15), (BODY, 11)],
            [("Sick Leave", 15), (BODY, 11)],
        ]
    )
    doc = parse_pdf(data, "fallback", S)
    assert doc.paged and doc.page_count == 2 and not doc.needs_ocr
    headings = [(b.text, b.level, b.page_start) for b in doc.blocks if b.kind == "heading"]
    assert headings == [("Annual Leave", 1, 1), ("Carry forward", 2, 1), ("Sick Leave", 2, 2)]
    paras = [b for b in doc.blocks if b.kind == "paragraph"]
    assert [p.page_start for p in paras] == [1, 1, 2]
    assert doc.title == "Annual Leave"


def test_pdf_drops_repeated_headers_and_page_numbers() -> None:
    pages = [[(f"Section {i}", 15), (f"{BODY} Unique {i}.", 11)] for i in range(4)]
    data = make_pdf(pages, header="ACME CONFIDENTIAL", footer_numbers=True)
    doc = parse_pdf(data, "f", S)
    texts = " ".join(b.text for b in doc.blocks)
    assert "ACME CONFIDENTIAL" not in texts
    assert "Page 1" not in texts and "Page 4" not in texts
    assert "Unique 2" in texts


def test_pdf_page_numbers_dropped_even_on_single_page() -> None:
    data = make_pdf([[("Title", 20), (BODY, 11)]], footer_numbers=True)
    assert "Page 1" not in " ".join(b.text for b in parse_pdf(data, "f", S).blocks)


def test_pdf_paragraph_continuing_across_page_break_is_merged() -> None:
    data = make_pdf(
        [
            [("Heading", 18), ("The retention period for shipment records is seven", 11)],
            [("years and applies to all depots.", 11)],
        ]
    )
    doc = parse_pdf(data, "f", S)
    (para,) = [b for b in doc.blocks if b.kind == "paragraph"]
    assert "seven years and applies" in para.text
    assert (para.page_start, para.page_end) == (1, 2)


def test_pdf_metadata_title_wins() -> None:
    data = make_pdf([[("Some Heading", 20), (BODY, 11)]], title="Real Title")
    assert parse_pdf(data, "f", S).title == "Real Title"


def test_scanned_pdf_is_flagged_needs_ocr_not_crashed() -> None:
    doc = parse_pdf(make_scanned_pdf(), "scan", S)
    assert doc.needs_ocr and doc.blocks == [] and doc.page_count == 1


def test_encrypted_pdf_is_a_clean_parse_error() -> None:
    with pytest.raises(ParseError, match="encrypted"):
        parse_pdf(make_encrypted_pdf(), "f", S)


def test_corrupt_pdf_is_a_clean_parse_error() -> None:
    with pytest.raises(ParseError):
        parse_pdf(b"%PDF-1.7 this is not really a pdf", "f", S)


def test_pdf_page_limit() -> None:
    data = make_pdf([[(BODY, 11)]] * 3)
    with pytest.raises(ParseError, match="too many pages"):
        parse_pdf(data, "f", Settings(pdf_max_pages=2))


# ---------------------------------------------------------------- DOCX
def test_docx_headings_paragraphs_and_tables() -> None:
    import io

    from docx import Document

    doc = Document()
    doc.add_heading("Policy", level=0)
    doc.add_heading("Leave", level=1)
    doc.add_heading("Carry forward", level=2)
    doc.add_paragraph("Ten days may be carried forward.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Type", "Days"
    table.cell(1, 0).text, table.cell(1, 1).text = "Sick", "12"
    buf = io.BytesIO()
    doc.save(buf)
    parsed = parse_docx(buf.getvalue(), "f", S)
    assert [(b.kind, b.level, b.text) for b in parsed.blocks] == [
        ("heading", 0, "Policy"),
        ("heading", 1, "Leave"),
        ("heading", 2, "Carry forward"),
        ("paragraph", 0, "Ten days may be carried forward."),
        ("paragraph", 0, "Type | Days"),
        ("paragraph", 0, "Sick | 12"),
    ]
    assert parsed.title == "Policy"


def test_docx_zip_bomb_rejected_before_inflating() -> None:
    base = make_docx([("A", ["b"])])
    bomb = rewrite_zip(base, {"word/padding.bin": b"\x00" * (60 * 1024 * 1024)})
    with pytest.raises(ParseError, match="compression ratio|too large"):
        parse_docx(bomb, "f", Settings(docx_max_compression_ratio=50.0))


def test_docx_too_many_entries_rejected() -> None:
    base = make_docx([("A", ["b"])])
    many = rewrite_zip(base, {f"word/x{i}.txt": b"x" for i in range(30)})
    with pytest.raises(ParseError, match="too many entries"):
        parse_docx(many, "f", Settings(docx_max_entries=10))


def test_docx_path_traversal_and_nested_archive_rejected() -> None:
    base = make_docx([("A", ["b"])])
    with pytest.raises(ParseError, match="unsafe path"):
        parse_docx(rewrite_zip(base, {"../evil.txt": b"x"}), "f", S)
    with pytest.raises(ParseError, match="nested archive"):
        parse_docx(rewrite_zip(base, {"word/embeddings/inner.zip": b"PK"}), "f", S)


def test_docx_xxe_and_entity_expansion_neutralised() -> None:
    base = make_docx([("A", ["b"])])
    xxe = (
        b'<?xml version="1.0"?><!DOCTYPE d [<!ENTITY x SYSTEM "file:///etc/passwd">]>'
        b'<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        b"<w:body><w:p><w:r><w:t>&x;</w:t></w:r></w:p></w:body></w:document>"
    )
    with pytest.raises(ParseError, match="forbidden XML"):
        parse_docx(rewrite_zip(base, {}, replace={"word/document.xml": xxe}), "f", S)
    laughs = (
        b'<?xml version="1.0"?><!DOCTYPE l [<!ENTITY a "aaaa"><!ENTITY b "&a;&a;&a;&a;">]>'
        b'<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        b"<w:body><w:p><w:r><w:t>&b;</w:t></w:r></w:p></w:body></w:document>"
    )
    with pytest.raises(ParseError, match="forbidden XML"):
        parse_docx(rewrite_zip(base, {}, replace={"word/document.xml": laughs}), "f", S)


def test_not_a_zip_is_a_clean_error() -> None:
    with pytest.raises(ParseError, match="bad zip"):
        parse_docx(b"definitely not a zip", "f", S)


# ---------------------------------------------------------------- HTML / text
HTML = """<html><head><title>Depot Guide</title></head><body>
<nav>Home About Contact us today</nav>
<article><h1>Depot Guide</h1>
<p>The depot opens at 06:00 and closes at 22:00 on weekdays, and the loading bay is
staffed throughout.</p>
<h2>Holidays</h2>
<p>Closed on public holidays except for emergency dispatch, which is handled by the duty
manager.</p>
<ul><li>Bring your badge every day</li><li>Sign in at the front desk on arrival</li></ul>
</article><footer>Copyright Northwind</footer>
<img src="http://evil.example/track.png"></body></html>"""


def test_html_extracts_main_content_and_heading_levels() -> None:
    doc = parse_html(HTML.encode(), "fallback")
    assert doc.title == "Depot Guide"
    heads = [(b.level, b.text) for b in doc.blocks if b.kind == "heading"]
    assert heads == [(1, "Depot Guide"), (2, "Holidays")]
    text = " ".join(b.text for b in doc.blocks)
    assert "opens at 06:00" in text and "Bring your badge" in text
    assert "Copyright" not in text and "Home About" not in text


def test_html_without_content_is_empty_not_error() -> None:
    assert parse_html(b"<html><body></body></html>", "x").blocks == []


def test_markdown_headings_and_paragraphs() -> None:
    md = b"# Guide\n\nIntro line one\ncontinues here.\n\n## Setup\n\nRun it.\n"
    doc = parse_text(md, "f", markdown=True)
    assert doc.title == "Guide"
    assert [(b.kind, b.level, b.text) for b in doc.blocks] == [
        ("heading", 1, "Guide"),
        ("paragraph", 0, "Intro line one continues here."),
        ("heading", 2, "Setup"),
        ("paragraph", 0, "Run it."),
    ]


def test_plain_text_has_no_headings() -> None:
    doc = parse_text(b"# not a heading\n\nsecond", "notes", markdown=False)
    assert all(b.kind == "paragraph" for b in doc.blocks) and doc.title == "notes"


def test_dispatch_rejects_unknown_mime() -> None:
    with pytest.raises(ParseError):
        parse_document("application/x-msdownload", b"MZ", "f", S)
    assert parse_document(MIME_PDF, make_pdf([[(BODY, 11)]]), "f", S).paged
