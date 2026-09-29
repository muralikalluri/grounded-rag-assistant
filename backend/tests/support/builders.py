"""Build small PDFs / DOCX files in memory for parser and pipeline tests."""

import io
import zipfile

import pymupdf
from docx import Document as DocxDocument

BODY = "Body text of the document that is long enough to count as content."


def make_pdf(
    pages: list[list[tuple[str, float]]],
    *,
    header: str | None = None,
    footer_numbers: bool = False,
    title: str = "",
) -> bytes:
    """pages = per page a list of (text, fontsize) drawn top to bottom."""
    doc = pymupdf.open()
    for n, items in enumerate(pages, start=1):
        page = doc.new_page(width=595, height=842)
        if header:
            page.insert_text((72, 35), header, fontsize=9)
        y = 100.0
        for text, size in items:
            page.insert_text((72, y), text, fontsize=size)
            y += size * 2.2
        if footer_numbers:
            page.insert_text((280, 815), f"Page {n}", fontsize=9)
    if title:
        doc.set_metadata({"title": title})
    data = doc.tobytes()
    doc.close()
    return data


def make_scanned_pdf() -> bytes:
    src = pymupdf.open()
    p = src.new_page()
    p.insert_text((72, 100), "Scanned text that is only an image", fontsize=14)
    pix = p.get_pixmap(dpi=40)
    out = pymupdf.open()
    page = out.new_page()
    page.insert_image(page.rect, pixmap=pix)
    data = out.tobytes()
    out.close()
    src.close()
    return data


def make_encrypted_pdf() -> bytes:
    doc = pymupdf.open()
    doc.new_page().insert_text((72, 100), "secret", fontsize=12)
    data = doc.tobytes(encryption=pymupdf.PDF_ENCRYPT_AES_256, owner_pw="owner", user_pw="user")
    doc.close()
    return data


def make_docx(sections: list[tuple[str, list[str]]], title: str = "Doc Title") -> bytes:
    doc = DocxDocument()
    doc.core_properties.title = title
    doc.add_heading(title, level=0)
    for heading, paras in sections:
        doc.add_heading(heading, level=1)
        for p in paras:
            doc.add_paragraph(p)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def rewrite_zip(
    original: bytes, extra: dict[str, bytes], *, replace: dict[str, bytes] | None = None
) -> bytes:
    buf = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(original)) as src,
        zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as out,
    ):
        for info in src.infolist():
            data = (replace or {}).get(info.filename, src.read(info.filename))
            out.writestr(info.filename, data)
        for name, data in extra.items():
            out.writestr(name, data)
    return buf.getvalue()


def make_page_pdf(
    pages: list[list[str]],
    *,
    rotation: int = 0,
    crop: tuple[float, float, float, float] | None = None,
) -> bytes:
    """One text line per entry, page by page (11pt), for locate/render tests."""
    doc = pymupdf.open()
    for lines in pages:
        page = doc.new_page(width=595, height=842)
        y = 100.0
        for line in lines:
            page.insert_text((72, y), line, fontsize=11)
            y += 16
        if crop:
            page.set_cropbox(pymupdf.Rect(*crop))
        if rotation:
            page.set_rotation(rotation)
    data = doc.tobytes()
    doc.close()
    return data
